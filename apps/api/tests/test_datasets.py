"""
Unit & Integration Tests untuk Dataset Endpoints dan Service (Tahap 2).
Mencakup seluruh 12 skenario validasi:
1. Impor CSV
2. Impor XLSX
3. File kosong dan file rusak
4. Kolom duplikat
5. Data numerik dan string
6. Missing values
7. Penambahan dan penghapusan baris
8. Penambahan dan penghapusan kolom
9. Perubahan Variable View
10. Penyimpanan dan pembukaan ulang dataset
11. Ekspor CSV dan XLSX
12. Konsistensi nilai sebelum dan sesudah ekspor
"""

import io
import os
import sys
import pytest
import pandas as pd
import openpyxl
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from main import app
from services.dataset_service import DatasetService

client = TestClient(app)


# ============================================================================
# 1 & 5. Impor CSV & Data Numerik/String
# ============================================================================

def test_import_csv_valid_numeric_and_string():
    """Menguji preview dan import file CSV dengan kolom numerik dan string."""
    csv_content = "ID,Nama,Skor,Status\n1,Budi,85.5,Lulus\n2,Siti,92.0,Lulus\n3,Joko,58.0,Remedial\n"
    files = {"file": ("test_students.csv", csv_content.encode("utf-8"), "text/csv")}

    # 1. Preview
    res_preview = client.post("/api/datasets/preview-import", files=files, data={"has_headers": True})
    assert res_preview.status_code == 200, res_preview.text
    preview = res_preview.json()
    assert preview["total_rows"] == 3
    assert preview["total_columns"] == 4
    col_types = {c["name"]: c["detected_type"] for c in preview["columns"]}
    assert col_types["ID"] == "numeric"
    assert col_types["Nama"] == "string"
    assert col_types["Skor"] == "numeric"
    assert col_types["Status"] == "string"

    # 2. Commit import
    variables = [
        {"name": c["name"], "type": c["detected_type"], "width": 8, "decimals": 2 if c["detected_type"] == "numeric" else 0, "label": c["name"], "measure": c["detected_measure"]}
        for c in preview["columns"]
    ]
    res_import = client.post(
        "/api/datasets/import",
        json={"name": "Siswa Test", "variables": variables, "rows": preview["preview_rows"]},
    )
    assert res_import.status_code == 200
    dataset = res_import.json()
    assert dataset["name"] == "Siswa Test"
    assert len(dataset["rows"]) == 3
    assert dataset["rows"][0]["Nama"] == "Budi"
    assert float(dataset["rows"][0]["Skor"]) == 85.5


# ============================================================================
# 2. Impor XLSX (Excel Multi-Sheet)
# ============================================================================

def test_import_xlsx_valid():
    """Menguji preview dan import file XLSX dengan pemilihan sheet."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Kuesioner"
    ws1.append(["Responden", "Q1", "Q2"])
    ws1.append([1, 4, 5])
    ws1.append([2, 3, 4])

    ws2 = wb.create_sheet(title="Info")
    ws2.append(["Keterangan", "Versi"])
    ws2.append(["Data Uji", "1.0"])

    excel_buf = io.BytesIO()
    wb.save(excel_buf)
    excel_buf.seek(0)

    files = {"file": ("survey.xlsx", excel_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res_preview = client.post("/api/datasets/preview-import", files=files, data={"sheet_name": "Kuesioner"})
    assert res_preview.status_code == 200
    preview = res_preview.json()
    assert "Kuesioner" in preview["sheet_names"]
    assert "Info" in preview["sheet_names"]
    assert preview["total_rows"] == 2
    assert preview["total_columns"] == 3


# ============================================================================
# 3. File Kosong dan File Rusak
# ============================================================================

def test_import_empty_file():
    """Menguji respons error terhadap file kosong (0 bytes)."""
    files = {"file": ("empty.csv", b"", "text/csv")}
    res = client.post("/api/datasets/preview-import", files=files)
    assert res.status_code == 400
    assert "kosong" in res.json()["detail"].lower()


def test_import_corrupted_file():
    """Menguji respons error terhadap file rusak."""
    files = {"file": ("corrupt.xlsx", b"Bukan file zip/excel valid", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res = client.post("/api/datasets/preview-import", files=files)
    assert res.status_code == 400
    assert "rusak" in res.json()["detail"].lower()


# ============================================================================
# 4. Kolom Duplikat
# ============================================================================

def test_duplicate_columns_handling():
    """Menguji penanganan nama kolom yang duplikat dengan renaming otomatis dan warning."""
    csv_content = "Umur,Umur,Skor,skor\n25,26,80,90\n30,31,85,95\n"
    files = {"file": ("duplicates.csv", csv_content.encode("utf-8"), "text/csv")}
    res = client.post("/api/datasets/preview-import", files=files)
    assert res.status_code == 200
    data = res.json()
    col_names = [c["name"] for c in data["columns"]]
    # Harus unik
    assert len(col_names) == len(set(col_names))
    assert any("_1" in col for col in col_names)
    assert len(data["warnings"]) > 0


# ============================================================================
# 6. Missing Values (System & User-Defined)
# ============================================================================

def test_missing_values_detection_and_count():
    """Menguji deteksi missing values kosong dan nilai kode missing khusus (misal 99)."""
    csv_content = "ID,Nilai\n1,10\n2,\n3,None\n4,99\n5,15\n"
    files = {"file": ("missing.csv", csv_content.encode("utf-8"), "text/csv")}
    res = client.post("/api/datasets/preview-import", files=files)
    assert res.status_code == 200
    preview = res.json()
    col_nilai = next(c for c in preview["columns"] if c["name"] == "Nilai")
    # baris 2 dan 3 adalah None/kosong -> terdeteksi missing
    assert col_nilai["missing_count"] >= 2


# ============================================================================
# 7. Penambahan dan Penghapusan Baris
# ============================================================================

def test_add_and_delete_rows():
    """Menguji API penambahan dan penghapusan baris."""
    # 1. Buat dataset
    res_create = client.post(
        "/api/datasets",
        json={"name": "Row Test Dataset", "variables": [{"name": "A", "type": "numeric"}], "rows": [{"A": 10}, {"A": 20}]},
    )
    assert res_create.status_code == 200
    ds_id = res_create.json()["id"]

    # 2. Tambah baris
    res_add = client.post(f"/api/datasets/{ds_id}/rows", json={"A": 30})
    assert res_add.status_code == 200
    rows = res_add.json()["rows"]
    assert len(rows) == 3
    assert rows[2]["A"] == 30

    # 3. Hapus baris indeks 1 (nilai 20)
    res_del = client.delete(f"/api/datasets/{ds_id}/rows/1")
    assert res_del.status_code == 200
    rows_after = res_del.json()["rows"]
    assert len(rows_after) == 2
    assert rows_after[0]["A"] == 10
    assert rows_after[1]["A"] == 30


# ============================================================================
# 8. Penambahan dan Penghapusan Kolom
# ============================================================================

def test_add_and_delete_columns():
    """Menguji API penambahan dan penghapusan kolom variabel."""
    res_create = client.post(
        "/api/datasets",
        json={"name": "Col Test", "variables": [{"name": "X", "type": "numeric"}], "rows": [{"X": 1}]},
    )
    ds_id = res_create.json()["id"]

    # Tambah kolom Y
    res_add_col = client.post(
        f"/api/datasets/{ds_id}/columns",
        json={"name": "Y", "type": "string", "label": "Kolom Y"},
    )
    assert res_add_col.status_code == 200
    vars_list = res_add_col.json()["variables"]
    assert any(v["name"] == "Y" for v in vars_list)
    assert "Y" in res_add_col.json()["rows"][0]

    # Hapus kolom X
    res_del_col = client.delete(f"/api/datasets/{ds_id}/columns/X")
    vars_after = res_del_col.json()["variables"]
    assert not any(v["name"] == "X" for v in vars_after)


# ============================================================================
# 9. Perubahan Variable View (SPSS Metadata)
# ============================================================================

def test_variable_view_update():
    """Menguji pengubahan definisi variabel: Value Labels, Decimals, Measure."""
    res_create = client.post(
        "/api/datasets",
        json={"name": "Variable View Test", "variables": [{"name": "Q1", "type": "numeric"}], "rows": [{"Q1": 1}]},
    )
    ds_id = res_create.json()["id"]

    # Update variabel Q1 dengan Value Labels kuesioner dan missing 99
    updated_vars = [
        {
            "name": "Q1",
            "type": "numeric",
            "width": 8,
            "decimals": 0,
            "label": "Tingkat Kepuasan",
            "values": {"1": "Puas", "2": "Tidak Puas"},
            "missing": ["99"],
            "measure": "ordinal",
        }
    ]
    res_update = client.put(f"/api/datasets/{ds_id}", json={"variables": updated_vars})
    assert res_update.status_code == 200
    v = res_update.json()["variables"][0]
    assert v["label"] == "Tingkat Kepuasan"
    assert v["values"]["1"] == "Puas"
    assert v["missing"] == ["99"]
    assert v["measure"] == "ordinal"


# ============================================================================
# 10. Penyimpanan dan Pembukaan Ulang Dataset (Persistence)
# ============================================================================

def test_save_and_reopen_dataset():
    """Menguji persistence: simpan dataset dan pastikan dapat dibaca kembali dengan data persis."""
    res_create = client.post(
        "/api/datasets",
        json={
            "name": "Persistent Dataset",
            "variables": [{"name": "Score", "type": "numeric"}],
            "rows": [{"Score": 99.5}, {"Score": 88.0}],
        },
    )
    assert res_create.status_code == 200
    ds_id = res_create.json()["id"]

    # Ambil ulang
    res_get = client.get(f"/api/datasets/{ds_id}")
    assert res_get.status_code == 200
    loaded = res_get.json()
    assert loaded["id"] == ds_id
    assert loaded["name"] == "Persistent Dataset"
    assert len(loaded["rows"]) == 2
    assert float(loaded["rows"][0]["Score"]) == 99.5


# ============================================================================
# 11 & 12. Ekspor CSV, XLSX, dan Konsistensi Nilai Sebelum & Sesudah Ekspor
# ============================================================================

def test_export_and_value_consistency():
    """Menguji ekspor ke CSV dan XLSX serta verifikasi konsistensi nilai."""
    test_rows = [
        {"id": 1, "ID": 101, "City": "Jakarta", "Income": 15000000.5},
        {"id": 2, "ID": 102, "City": "Surabaya", "Income": 12500000.0},
        {"id": 3, "ID": 103, "City": "Bandung", "Income": None},
    ]
    test_vars = [
        {"name": "ID", "type": "numeric", "label": "Nomor Karyawan", "measure": "nominal"},
        {"name": "City", "type": "string", "label": "Kota", "measure": "nominal"},
        {"name": "Income", "type": "numeric", "decimals": 2, "label": "Pendapatan Bulanan", "measure": "scale"},
    ]

    res_create = client.post(
        "/api/datasets",
        json={"name": "Export Test", "variables": test_vars, "rows": test_rows},
    )
    ds_id = res_create.json()["id"]

    # 1. Ekspor CSV
    res_csv = client.get(f"/api/datasets/{ds_id}/export/csv")
    assert res_csv.status_code == 200
    assert "text/csv" in res_csv.headers["content-type"]
    csv_text = res_csv.text

    # Baca ulang CSV dengan pandas dan periksa konsistensi
    df_csv = pd.read_csv(io.StringIO(csv_text))
    assert list(df_csv.columns) == ["ID", "City", "Income"]
    assert df_csv.iloc[0]["City"] == "Jakarta"
    assert float(df_csv.iloc[0]["Income"]) == 15000000.5
    assert pd.isna(df_csv.iloc[2]["Income"])

    # 2. Ekspor XLSX
    res_xlsx = client.get(f"/api/datasets/{ds_id}/export/xlsx")
    assert res_xlsx.status_code == 200
    assert "openxmlformats" in res_xlsx.headers["content-type"]

    # Baca ulang XLSX dengan openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    assert "Data" in wb.sheetnames
    assert "Variable_Dictionary" in wb.sheetnames
    ws_data = wb["Data"]
    assert ws_data.cell(row=1, column=1).value == "ID"
    assert ws_data.cell(row=2, column=2).value == "Jakarta"
    assert float(ws_data.cell(row=2, column=3).value) == 15000000.5


# ============================================================================
# DEMO DATA Seed & Reset
# ============================================================================

def test_demo_data_seeding_and_structure():
    """Menguji bahwa DEMO DATA di-generate dengan variabel ID, Q1, Q2, Q3, Q4 dan simulasi 25 responden."""
    res_seed = client.post("/api/datasets/demo/seed?force=true")
    assert res_seed.status_code == 200
    demo = res_seed.json()
    assert demo["is_demo"] is True
    var_names = [v["name"] for v in demo["variables"]]
    assert var_names == ["ID", "Q1", "Q2", "Q3", "Q4"]
    assert len(demo["rows"]) == 25
    # Periksa value labels pada Q1
    q1 = next(v for v in demo["variables"] if v["name"] == "Q1")
    assert q1["values"]["1"] == "Sangat Tidak Setuju"
    assert q1["values"]["5"] == "Sangat Setuju"
