import os
import json
import uuid
import io
import re
from datetime import datetime, timezone
from typing import List, Dict, Any, Optional, Tuple
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from config import settings
from services.storage_security import atomic_write_json, sanitize_formula_cell
from schemas.dataset import (
    VariableDefinition,
    DatasetMetadata,
    Dataset,
    DatasetCreate,
    DatasetUpdate,
    ColumnPreview,
    ImportPreviewResponse,
    ImportCommitRequest,
    sanitize_var_name,
)

DATA_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), settings.data_dir, "datasets")
INDEX_FILE = os.path.join(DATA_DIR, "index.json")


def _ensure_storage_dir():
    os.makedirs(DATA_DIR, exist_ok=True)
    if not os.path.exists(INDEX_FILE):
        with open(INDEX_FILE, "w", encoding="utf-8") as f:
            json.dump([], f)


def _load_index() -> List[Dict[str, Any]]:
    _ensure_storage_dir()
    try:
        with open(INDEX_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return []


def _save_index(items: List[Dict[str, Any]]) -> None:
    _ensure_storage_dir()
    atomic_write_json(INDEX_FILE, items, indent=2)


def _calculate_missing(variables: List[VariableDefinition], rows: List[Dict[str, Any]]) -> int:
    """Hitung total missing values di seluruh sel."""
    missing_lookup = {v.name: set(v.missing) for v in variables}
    count = 0
    for row in rows:
        for v in variables:
            val = row.get(v.name)
            if val is None or val == "" or str(val).strip() == "" or str(val).lower() in ("nan", "none", "null"):
                count += 1
            elif str(val) in missing_lookup.get(v.name, set()):
                count += 1
    return count


def _to_native(v: Any) -> Any:
    if v is None or pd.isna(v):
        return None
    if hasattr(v, "item"):
        return v.item()
    return v


class DatasetService:
    @staticmethod
    def list_datasets() -> List[DatasetMetadata]:
        _ensure_storage_dir()
        index = _load_index()
        # Jika belum ada dataset, inisialisasi DEMO DATA otomatis
        if not index:
            demo = DatasetService.seed_demo_data()
            return [
                DatasetMetadata(
                    id=demo.id,
                    name=demo.name,
                    description=demo.description,
                    row_count=len(demo.rows),
                    variable_count=len(demo.variables),
                    missing_count=_calculate_missing(demo.variables, demo.rows),
                    is_demo=demo.is_demo,
                    created_at=demo.created_at,
                    updated_at=demo.updated_at,
                )
            ]

        result = []
        for item in index:
            result.append(
                DatasetMetadata(
                    id=item["id"],
                    name=item["name"],
                    description=item.get("description", ""),
                    row_count=item.get("row_count", 0),
                    variable_count=item.get("variable_count", 0),
                    missing_count=item.get("missing_count", 0),
                    is_demo=item.get("is_demo", False),
                    created_at=item.get("created_at", ""),
                    updated_at=item.get("updated_at", ""),
                )
            )
        return result

    @staticmethod
    def get_dataset(dataset_id: str) -> Optional[Dataset]:
        _ensure_storage_dir()
        file_path = os.path.join(DATA_DIR, f"{dataset_id}.json")
        if not os.path.exists(file_path):
            return None
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            return Dataset(**data)
        except Exception:
            return None

    @staticmethod
    def create_dataset(data: DatasetCreate, is_demo: bool = False, dataset_id: Optional[str] = None) -> Dataset:
        _ensure_storage_dir()
        ds_id = dataset_id or str(uuid.uuid4())
        now = datetime.now(timezone.utc).isoformat()

        variables = data.variables if data.variables is not None else [
            VariableDefinition(name="VAR00001", type="numeric", width=8, decimals=2, label="", measure="scale"),
            VariableDefinition(name="VAR00002", type="numeric", width=8, decimals=2, label="", measure="scale"),
            VariableDefinition(name="VAR00003", type="numeric", width=8, decimals=2, label="", measure="scale"),
        ]

        rows = data.rows or []
        for i, row in enumerate(rows):
            if "id" not in row:
                row["id"] = i + 1

        dataset = Dataset(
            id=ds_id,
            name=data.name.strip() or "Dataset Baru",
            description=data.description or "",
            variables=variables,
            rows=rows,
            is_demo=is_demo,
            created_at=now,
            updated_at=now,
        )

        # Simpan file dataset secara atomik
        file_path = os.path.join(DATA_DIR, f"{ds_id}.json")
        atomic_write_json(file_path, dataset.model_dump(), indent=2)

        # Update index
        index = _load_index()
        meta = {
            "id": ds_id,
            "name": dataset.name,
            "description": dataset.description,
            "row_count": len(dataset.rows),
            "variable_count": len(dataset.variables),
            "missing_count": _calculate_missing(dataset.variables, dataset.rows),
            "is_demo": is_demo,
            "created_at": now,
            "updated_at": now,
        }
        index = [item for item in index if item["id"] != ds_id]
        index.insert(0, meta)
        _save_index(index)

        return dataset

    @staticmethod
    def update_dataset(dataset_id: str, update_data: DatasetUpdate) -> Optional[Dataset]:
        _ensure_storage_dir()
        existing = DatasetService.get_dataset(dataset_id)
        if not existing:
            return None

        now = datetime.now(timezone.utc).isoformat()
        if update_data.name is not None and update_data.name.strip():
            existing.name = update_data.name.strip()
        if update_data.description is not None:
            existing.description = update_data.description
        if update_data.variables is not None:
            # Validasi nama duplikat antar variabel
            seen = set()
            for v in update_data.variables:
                if v.name.lower() in seen:
                    raise ValueError(f"Nama variabel duplikat terdeteksi: {v.name}")
                seen.add(v.name.lower())
            existing.variables = update_data.variables
        if update_data.rows is not None:
            existing.rows = update_data.rows

        existing.updated_at = now

        file_path = os.path.join(DATA_DIR, f"{dataset_id}.json")
        atomic_write_json(file_path, existing.model_dump(), indent=2)

        # Update index
        index = _load_index()
        for item in index:
            if item["id"] == dataset_id:
                item["name"] = existing.name
                item["description"] = existing.description
                item["row_count"] = len(existing.rows)
                item["variable_count"] = len(existing.variables)
                item["missing_count"] = _calculate_missing(existing.variables, existing.rows)
                item["updated_at"] = now
                break
        _save_index(index)

        return existing

    @staticmethod
    def delete_dataset(dataset_id: str) -> bool:
        _ensure_storage_dir()
        file_path = os.path.join(DATA_DIR, f"{dataset_id}.json")
        existed = False
        if os.path.exists(file_path):
            os.remove(file_path)
            existed = True

        index = _load_index()
        new_index = [item for item in index if item["id"] != dataset_id]
        if len(new_index) != len(index):
            existed = True
            _save_index(new_index)

        return existed

    @staticmethod
    def add_row(dataset_id: str, row_data: Dict[str, Any]) -> Optional[Dataset]:
        ds = DatasetService.get_dataset(dataset_id)
        if not ds:
            return None
        new_id = (max([r.get("id", 0) for r in ds.rows], default=0) + 1)
        row_data["id"] = new_id
        # Inisialisasi variabel yang belum ada
        for v in ds.variables:
            if v.name not in row_data:
                row_data[v.name] = None
        ds.rows.append(row_data)
        return DatasetService.update_dataset(dataset_id, DatasetUpdate(rows=ds.rows))

    @staticmethod
    def delete_row(dataset_id: str, row_index: int) -> Optional[Dataset]:
        ds = DatasetService.get_dataset(dataset_id)
        if not ds:
            return None
        if row_index < 0 or row_index >= len(ds.rows):
            raise IndexError("Indeks baris berada di luar batas")
        ds.rows.pop(row_index)
        return DatasetService.update_dataset(dataset_id, DatasetUpdate(rows=ds.rows))

    @staticmethod
    def add_column(dataset_id: str, var_def: VariableDefinition) -> Optional[Dataset]:
        ds = DatasetService.get_dataset(dataset_id)
        if not ds:
            return None
        if any(v.name.lower() == var_def.name.lower() for v in ds.variables):
            raise ValueError(f"Variabel '{var_def.name}' sudah ada")
        ds.variables.append(var_def)
        for r in ds.rows:
            r[var_def.name] = None
        return DatasetService.update_dataset(dataset_id, DatasetUpdate(variables=ds.variables, rows=ds.rows))

    @staticmethod
    def delete_column(dataset_id: str, var_name: str) -> Optional[Dataset]:
        ds = DatasetService.get_dataset(dataset_id)
        if not ds:
            return None
        if not any(v.name.lower() == var_name.lower() for v in ds.variables):
            raise ValueError(f"Variabel '{var_name}' tidak ditemukan")
        ds.variables = [v for v in ds.variables if v.name.lower() != var_name.lower()]
        for r in ds.rows:
            r.pop(var_name, None)
        return DatasetService.update_dataset(dataset_id, DatasetUpdate(variables=ds.variables, rows=ds.rows))

    @staticmethod
    def seed_demo_data(force: bool = False) -> Dataset:
        _ensure_storage_dir()
        DEMO_ID = "demo-dataset-spss"
        existing = DatasetService.get_dataset(DEMO_ID)
        if existing and not force:
            return existing

        variables = [
            VariableDefinition(
                name="ID",
                type="numeric",
                width=4,
                decimals=0,
                label="Nomor Responden (DEMO)",
                measure="nominal",
            ),
            VariableDefinition(
                name="Q1",
                type="numeric",
                width=4,
                decimals=0,
                label="Transparansi Pengelolaan Anggaran",
                values={
                    "1": "Sangat Tidak Setuju",
                    "2": "Tidak Setuju",
                    "3": "Netral",
                    "4": "Setuju",
                    "5": "Sangat Setuju",
                },
                missing=["99"],
                measure="ordinal",
            ),
            VariableDefinition(
                name="Q2",
                type="numeric",
                width=4,
                decimals=0,
                label="Akuntabilitas Pelaporan Kinerja",
                values={
                    "1": "Sangat Tidak Setuju",
                    "2": "Tidak Setuju",
                    "3": "Netral",
                    "4": "Setuju",
                    "5": "Sangat Setuju",
                },
                missing=["99"],
                measure="ordinal",
            ),
            VariableDefinition(
                name="Q3",
                type="numeric",
                width=4,
                decimals=0,
                label="Kualitas Pelayanan Publik",
                values={
                    "1": "Sangat Rendah",
                    "2": "Rendah",
                    "3": "Cukup",
                    "4": "Tinggi",
                    "5": "Sangat Tinggi",
                },
                missing=["99"],
                measure="ordinal",
            ),
            VariableDefinition(
                name="Q4",
                type="numeric",
                width=4,
                decimals=0,
                label="Kepuasan Pengguna Layanan",
                values={
                    "1": "Sangat Tidak Puas",
                    "2": "Tidak Puas",
                    "3": "Cukup Puas",
                    "4": "Puas",
                    "5": "Sangat Puas",
                },
                missing=["99"],
                measure="ordinal",
            ),
        ]

        # 25 Baris Data Simulasi Demo Kuesioner (Jelas berlabel DEMO)
        demo_rows_raw = [
            (1, 4, 5, 4, 5),
            (2, 5, 4, 5, 4),
            (3, 3, 3, 4, 3),
            (4, 2, 2, 3, 2),
            (5, 4, 4, 4, 4),
            (6, 5, 5, 5, 5),
            (7, 1, 2, 2, 1),
            (8, 4, 3, 4, 4),
            (9, 3, 4, 3, 3),
            (10, 5, 5, 4, 5),
            (11, 4, 4, 5, 4),
            (12, 2, 3, 2, 3),
            (13, 3, 2, 3, 2),
            (14, 4, 5, 4, 4),
            (15, 5, 4, 5, 5),
            (16, 99, 4, 3, 4),  # Simulates user-defined missing (99)
            (17, 3, 3, 3, 3),
            (18, 4, 4, 4, 5),
            (19, 5, 5, 5, 4),
            (20, 2, 1, 2, 2),
            (21, 4, 4, 4, 4),
            (22, 3, 4, 4, 3),
            (23, 5, 5, 5, 5),
            (24, 4, 3, 4, 4),
            (25, None, 4, 4, 4),  # Simulates system-missing (None)
        ]

        rows = []
        for r in demo_rows_raw:
            rows.append({
                "id": r[0],
                "ID": r[0],
                "Q1": r[1],
                "Q2": r[2],
                "Q3": r[3],
                "Q4": r[4],
            })

        create_data = DatasetCreate(
            name="DEMO DATA — Simulasi Kuesioner",
            description="Dataset simulasi untuk pengujian fitur SPSS (bukan data responden asli).",
            variables=variables,
            rows=rows,
        )
        return DatasetService.create_dataset(create_data, is_demo=True, dataset_id=DEMO_ID)

    # ========================================================
    # IMPORT LOGIC (Pandas + OpenPyXL)
    # ========================================================

    @staticmethod
    def preview_import_file(
        content: bytes,
        filename: str,
        sheet_name: Optional[str] = None,
        delimiter: Optional[str] = None,
        has_headers: bool = True,
    ) -> ImportPreviewResponse:
        """Parse file dan generate preview lengkap dengan deteksi tipe data."""
        if not content:
            raise ValueError("File kosong (0 bytes). Harap pilih file yang valid.")

        max_bytes = settings.upload_max_size_mb * 1024 * 1024
        if len(content) > max_bytes:
            raise ValueError(f"Ukuran file melebihi batas maksimum {settings.upload_max_size_mb} MB.")

        ext = os.path.splitext(filename)[1].lower()
        if ext not in (".csv", ".xlsx", ".xls"):
            raise ValueError(f"Format file '{ext}' tidak didukung. Harap gunakan file CSV atau Excel (.xlsx/.xls).")

        warnings = []
        sheet_names = []
        selected_sheet = None
        df: pd.DataFrame

        if ext in (".xlsx", ".xls"):
            file_format = "xlsx"
            excel_file = io.BytesIO(content)
            try:
                wb = openpyxl.load_workbook(excel_file, read_only=True, data_only=True)
                sheet_names = wb.sheetnames
                wb.close()
            except Exception as e:
                raise ValueError(f"File Excel rusak atau tidak dapat dibaca: {str(e)}")

            if not sheet_names:
                raise ValueError("File Excel tidak memiliki lembar kerja (worksheet).")

            selected_sheet = sheet_name if sheet_name and sheet_name in sheet_names else sheet_names[0]
            excel_file.seek(0)
            try:
                header_arg = 0 if has_headers else None
                df = pd.read_excel(excel_file, sheet_name=selected_sheet, header=header_arg)
            except Exception as e:
                raise ValueError(f"Gagal membaca sheet '{selected_sheet}': {str(e)}")
        else:
            file_format = "csv"
            # Deteksi encoding
            text_data = None
            for enc in ("utf-8-sig", "utf-8", "latin1", "cp1252"):
                try:
                    text_data = content.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue

            if text_data is None:
                raise ValueError("Gagal mendeteksi encoding teks file CSV.")

            # Sniff delimiter jika tidak diberikan
            detected_delimiter = delimiter
            if not detected_delimiter:
                first_lines = "\n".join(text_data.splitlines()[:5])
                candidates = [",", ";", "\t", "|"]
                counts = {c: first_lines.count(c) for c in candidates}
                detected_delimiter = max(counts, key=counts.get) if any(counts.values()) else ","

            header_arg = 0 if has_headers else None
            try:
                df = pd.read_csv(
                    io.StringIO(text_data),
                    sep=detected_delimiter,
                    header=header_arg,
                    engine="python",
                )
            except Exception as e:
                raise ValueError(f"Gagal mem-parsing file CSV: {str(e)}")

        if df.empty and len(df.columns) == 0:
            raise ValueError("File tidak memuat data baris maupun kolom.")

        # Handle header nama kolom dan deteksi duplikasi
        raw_cols = list(df.columns)
        unique_cols: List[str] = []
        seen_names: Dict[str, int] = {}
        for i, col in enumerate(raw_cols):
            base_name = str(col).strip() if has_headers else f"VAR{str(i+1).zfill(5)}"
            clean_name = sanitize_var_name(base_name)
            lower_name = clean_name.lower()
            if lower_name in seen_names:
                seen_names[lower_name] += 1
                unique_name = f"{clean_name}_{seen_names[lower_name]}"
                warnings.append(f"Kolom duplikat '{col}' otomatis diubah menjadi '{unique_name}'.")
            else:
                seen_names[lower_name] = 0
                unique_name = clean_name
            unique_cols.append(unique_name)

        df.columns = unique_cols

        # Analisis tipe data kolom & missing
        columns_preview: List[ColumnPreview] = []
        for col in unique_cols:
            series = df[col]
            missing_count = int(series.isna().sum())
            non_null = series.dropna()

            # Deteksi tipe numeric vs string
            is_numeric = False
            if len(non_null) > 0:
                try:
                    pd.to_numeric(non_null)
                    is_numeric = True
                except (ValueError, TypeError):
                    is_numeric = False
            else:
                is_numeric = True  # default numeric jika semua kosong

            col_type = "numeric" if is_numeric else "string"
            measure = "scale" if is_numeric else "nominal"

            # Sampel nilai (maks 5)
            sample = [_to_native(v) for v in series.head(5).tolist()]

            columns_preview.append(
                ColumnPreview(
                    name=col,
                    original_name=str(raw_cols[len(columns_preview)]),
                    detected_type=col_type,
                    detected_measure=measure,
                    missing_count=missing_count,
                    sample_values=sample,
                )
            )

        # Generate preview rows (15 baris pertama)
        preview_df = df.head(15).copy()
        # Bersihkan NaN dan tipe numpy menjadi native python untuk JSON compatibility
        preview_rows = []
        for i, row in preview_df.iterrows():
            row_dict = {"id": i + 1}
            for col in unique_cols:
                row_dict[col] = _to_native(row[col])
            preview_rows.append(row_dict)

        return ImportPreviewResponse(
            filename=filename,
            file_format=file_format,
            sheet_names=sheet_names,
            selected_sheet=selected_sheet,
            delimiter=detected_delimiter if file_format == "csv" else None,
            has_headers=has_headers,
            total_rows=len(df),
            total_columns=len(unique_cols),
            columns=columns_preview,
            preview_rows=preview_rows,
            warnings=warnings,
        )

    @staticmethod
    def commit_import(request: ImportCommitRequest) -> Dataset:
        """Simpan data yang diimpor sebagai dataset baru atau ganti dataset yang ditentukan."""
        # Sanitize values dalam rows
        sanitized_rows = []
        var_types = {v.name: v.type for v in request.variables}

        for idx, row in enumerate(request.rows):
            new_row = {"id": row.get("id", idx + 1)}
            for v in request.variables:
                val = row.get(v.name)
                if val is None or val == "" or str(val).lower() in ("nan", "none", "null"):
                    new_row[v.name] = None
                elif var_types.get(v.name) == "numeric":
                    try:
                        new_row[v.name] = float(val) if "." in str(val) else int(val)
                    except (ValueError, TypeError):
                        new_row[v.name] = str(val)
                else:
                    new_row[v.name] = str(val)
            sanitized_rows.append(new_row)

        if request.target_dataset_id:
            # Gantikan dataset yang ada
            existing = DatasetService.get_dataset(request.target_dataset_id)
            if existing:
                update_dto = DatasetUpdate(
                    name=request.name or existing.name,
                    variables=request.variables,
                    rows=sanitized_rows,
                )
                return DatasetService.update_dataset(request.target_dataset_id, update_dto)

        # Buat dataset baru
        create_dto = DatasetCreate(
            name=request.name or "Dataset Hasil Impor",
            description=f"Diimpor pada {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')}",
            variables=request.variables,
            rows=sanitized_rows,
        )
        return DatasetService.create_dataset(create_dto)

    # ========================================================
    # EXPORT LOGIC (CSV & Multi-sheet XLSX)
    # ========================================================

    @staticmethod
    def export_csv(dataset_id: str) -> str:
        ds = DatasetService.get_dataset(dataset_id)
        if not ds:
            raise ValueError(f"Dataset {dataset_id} tidak ditemukan")

        col_names = [v.name for v in ds.variables]
        df_rows = []
        for r in ds.rows:
            row_data = {c: sanitize_formula_cell(r.get(c, "")) for c in col_names}
            df_rows.append(row_data)

        df = pd.DataFrame(df_rows, columns=col_names)
        output = io.StringIO()
        df.to_csv(output, index=False, encoding="utf-8")
        return output.getvalue()

    @staticmethod
    def export_xlsx(dataset_id: str) -> io.BytesIO:
        ds = DatasetService.get_dataset(dataset_id)
        if not ds:
            raise ValueError(f"Dataset {dataset_id} tidak ditemukan")

        wb = openpyxl.Workbook()

        # Sheet 1: Data
        ws_data = wb.active
        ws_data.title = "Data"

        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        border_thin = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin)

        # Tulis Headers Sheet Data
        col_names = [v.name for v in ds.variables]
        for col_idx, col_name in enumerate(col_names, 1):
            cell = ws_data.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Tulis Rows Sheet Data
        for row_idx, r in enumerate(ds.rows, 2):
            for col_idx, col_name in enumerate(col_names, 1):
                val = r.get(col_name)
                cell = ws_data.cell(row=row_idx, column=col_idx, value=val)
                cell.border = cell_border
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right")

        # Sheet 2: Variable Dictionary
        ws_dict = wb.create_sheet(title="Variable_Dictionary")
        dict_headers = ["Name", "Type", "Width", "Decimals", "Label", "Values", "Missing", "Measure"]
        for col_idx, h in enumerate(dict_headers, 1):
            cell = ws_dict.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx, v in enumerate(ds.variables, 2):
            values_str = "; ".join([f"{k} = {val}" for k, val in v.values.items()]) if v.values else "None"
            missing_str = ", ".join(v.missing) if v.missing else "None"
            row_vals = [
                v.name,
                v.type,
                v.width,
                v.decimals,
                v.label,
                values_str,
                missing_str,
                v.measure,
            ]
            for col_idx, val in enumerate(row_vals, 1):
                cell = ws_dict.cell(row=row_idx, column=col_idx, value=val)
                cell.border = cell_border

        # Auto fit column widths
        for ws in (ws_data, ws_dict):
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
