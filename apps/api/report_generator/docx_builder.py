import io
from typing import Dict, Any, List
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Mengatur padding margin sel Word."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def _apply_apa_table_borders(table):
    """Menerapkan garis batas tabel bergaya APA 7: hanya garis horizontal atas, bawah header, dan bawah tabel."""
    tblPr = table._tbl.tblPr
    # Hapus border bawaan dan terapkan garis horizontal APA 7
    borders_xml = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="12" w:space="0" w:color="0F172A"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="0F172A"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_xml)

    # Tambahkan garis bawah pada baris pertama (header)
    if len(table.rows) > 0:
        for cell in table.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            b_xml = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="0" w:color="0F172A"/></w:tcBorders>')
            tcPr.append(b_xml)


def build_docx_report(report_data: Dict[str, Any]) -> io.BytesIO:
    """Membangun dokumen Word DOCX berstandar APA 7 menggunakan python-docx."""
    doc = docx.Document()

    # Set margins 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    meta = report_data.get("metadata", {})
    ds = report_data.get("dataset", {})
    sel_sections = set(report_data.get("selected_sections", []))
    analysis_map = report_data.get("analysis_map", {})
    histories = report_data.get("histories", [])

    # 1. JUDUL & INFORMASI PENELITIAN
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run(meta.get("title") or "Laporan Hasil Analisis Statistik")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Statistika Pro — Applied Statistical Analysis System")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # Kotak Ringkasan Metadata
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_items = [
        ("Peneliti", meta.get("researcher") or "-"),
        ("Institusi / Jurusan", meta.get("institution") or "-"),
        ("Periode Penelitian", meta.get("period") or "-"),
        ("Nama Dataset", ds.get("name") or "-"),
        ("Ukuran Sampel (N)", f"{len(ds.get('rows', []))} Responden / Kasus"),
        ("Jumlah Variabel", f"{len(ds.get('variables', []))} Kolom"),
        ("Tanggal Dihasilkan", datetime.now().strftime("%d %B %Y %H:%M WIB")),
    ]
    if meta.get("notes"):
        meta_items.append(("Catatan", meta.get("notes")))

    for label, val in meta_items:
        row = meta_table.add_row()
        c0, c1 = row.cells
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        r0 = c0.paragraphs[0].add_run(label + ":")
        r0.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(71, 85, 105)
        r1 = c1.paragraphs[0].add_run(val)
        r1.font.size = Pt(10)
        _set_cell_margins(c0, 60, 60, 100, 100)
        _set_cell_margins(c1, 60, 60, 100, 100)

    _apply_apa_table_borders(meta_table)

    doc.add_page_break()

    # 2. KAMUS VARIABEL
    if "kamus_variabel" in sel_sections and ds.get("variables"):
        h1 = doc.add_heading("1. Kamus Variabel (Codebook)", level=1)
        h1.runs[0].font.color.rgb = RGBColor(30, 58, 138)
        p_desc = doc.add_paragraph(
            "Tabel berikut merangkum definisi variabel penelitian, tipe data, skala pengukuran, dan label nilai dalam dataset."
        )
        p_desc.runs[0].font.size = Pt(10)

        var_table = doc.add_table(rows=1, cols=6)
        var_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["No", "Nama", "Label Variabel", "Tipe", "Skala Ukur", "Missing"]
        for idx, h_text in enumerate(headers):
            cell = var_table.rows[0].cells[idx]
            run = cell.paragraphs[0].add_run(h_text)
            run.bold = True
            run.font.size = Pt(9.5)
            _set_cell_margins(cell, 80, 80, 100, 100)

        for idx, v in enumerate(ds.get("variables", [])):
            row = var_table.add_row()
            missing_str = ", ".join(str(m) for m in v.get("missing", [])) or "None"
            vals = [
                str(idx + 1),
                v.get("name", ""),
                v.get("label", "") or "-",
                v.get("type", "numeric"),
                v.get("measure", "scale"),
                missing_str,
            ]
            for c_idx, val_str in enumerate(vals):
                c = row.cells[c_idx]
                r = c.paragraphs[0].add_run(val_str)
                r.font.size = Pt(9)
                _set_cell_margins(c, 60, 60, 100, 100)

        _apply_apa_table_borders(var_table)
        doc.add_paragraph()

    # 3. RIWAYAT TRANSFORMASI
    if "riwayat_transformasi" in sel_sections and histories:
        h2 = doc.add_heading("2. Riwayat Transformasi Data (Audit Trail)", level=1)
        h2.runs[0].font.color.rgb = RGBColor(30, 58, 138)
        doc.add_paragraph("Rekaman audit operasi transformasi pada dataset aktif:").runs[0].font.size = Pt(10)

        hist_table = doc.add_table(rows=1, cols=4)
        hist_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        h_cols = ["Waktu (UTC)", "Operasi", "Keterangan", "Baris Terdampak"]
        for idx, col_name in enumerate(h_cols):
            cell = hist_table.rows[0].cells[idx]
            run = cell.paragraphs[0].add_run(col_name)
            run.bold = True
            run.font.size = Pt(9.5)
            _set_cell_margins(cell, 80, 80, 100, 100)

        for h in histories[:15]:
            row = hist_table.add_row()
            ts_clean = h.get("timestamp", "").replace("T", " ")[:19]
            h_vals = [ts_clean, h.get("operation_type", "-"), h.get("description", "-"), str(h.get("affected_rows", "-"))]
            for c_idx, val_str in enumerate(h_vals):
                c = row.cells[c_idx]
                r = c.paragraphs[0].add_run(val_str)
                r.font.size = Pt(9)
                _set_cell_margins(c, 60, 60, 100, 100)

        _apply_apa_table_borders(hist_table)
        doc.add_paragraph()

    # 4. HASIL ANALISIS STATISTIK
    section_mapping = [
        ("deskriptif", "descriptive", "Statistik Deskriptif"),
        ("reliabilitas", "reliability", "Uji Reliabilitas Instrumen"),
        ("validitas", "validity", "Uji Validitas Item"),
        ("korelasi", "correlation", "Analisis Korelasi"),
        ("regresi", "regression", "Analisis Regresi Linier"),
        ("diagnostik", "diagnostics", "Uji Asumsi Klasik"),
    ]

    sec_counter = 3
    for sec_key, anal_type, sec_title in section_mapping:
        if sec_key not in sel_sections:
            continue
        outputs_list = analysis_map.get(anal_type, [])
        if not outputs_list:
            continue

        h_sec = doc.add_heading(f"{sec_counter}. {sec_title}", level=1)
        h_sec.runs[0].font.color.rgb = RGBColor(30, 58, 138)
        sec_counter += 1

        for out in outputs_list:
            h_tbl = doc.add_heading(f"Tabel: {out.title}", level=2)
            h_tbl.runs[0].font.color.rgb = RGBColor(51, 65, 85)

            if out.dataset_updated_at and ds.get("updated_at") and out.dataset_updated_at != ds.get("updated_at"):
                p_warn = doc.add_paragraph("Perhatian: Dataset telah diperbarui sejak analisis ini dijalankan.")
                p_warn.runs[0].italic = True
                p_warn.runs[0].font.color.rgb = RGBColor(217, 119, 6)

            res = out.results
            if isinstance(res, dict):
                tables_to_draw = []
                if "tables" in res and isinstance(res["tables"], list):
                    tables_to_draw = res["tables"]
                elif "statistics" in res and isinstance(res["statistics"], list):
                    tables_to_draw = [{"title": "Ringkasan Statistik", "rows": res["statistics"]}]

                if tables_to_draw:
                    for t in tables_to_draw:
                        t_rows_raw = t.get("rows", [])
                        if not t_rows_raw:
                            continue
                        cols = list(t_rows_raw[0].keys())
                        stat_table = doc.add_table(rows=1, cols=len(cols))
                        stat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                        for c_i, c_name in enumerate(cols):
                            cell = stat_table.rows[0].cells[c_i]
                            r = cell.paragraphs[0].add_run(str(c_name))
                            r.bold = True
                            r.font.size = Pt(9)
                            _set_cell_margins(cell, 80, 80, 80, 80)

                        for r_item in t_rows_raw:
                            row = stat_table.add_row()
                            for c_i, c_name in enumerate(cols):
                                cell = row.cells[c_i]
                                val = r_item.get(c_name)
                                val_str = f"{val:.4f}" if isinstance(val, float) else str(val if val is not None else "-")
                                r = cell.paragraphs[0].add_run(val_str)
                                r.font.size = Pt(8.5)
                                _set_cell_margins(cell, 50, 50, 80, 80)

                        _apply_apa_table_borders(stat_table)
                        doc.add_paragraph()
                else:
                    # Key-value summary
                    kv_table = doc.add_table(rows=0, cols=2)
                    kv_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for k, v in list(res.items())[:12]:
                        if isinstance(v, (int, float, str, bool)):
                            row = kv_table.add_row()
                            c0, c1 = row.cells
                            r0 = c0.paragraphs[0].add_run(str(k))
                            r0.bold = True
                            r0.font.size = Pt(9)
                            r1 = c1.paragraphs[0].add_run(str(v))
                            r1.font.size = Pt(9)
                    _apply_apa_table_borders(kv_table)
                    doc.add_paragraph()

            if out.interpretation:
                p_interp = doc.add_paragraph()
                r_label = p_interp.add_run("Interpretasi: ")
                r_label.bold = True
                p_interp.add_run(out.interpretation)

            doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
