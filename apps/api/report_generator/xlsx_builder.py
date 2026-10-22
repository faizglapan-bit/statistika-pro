import io
from typing import Dict, Any, List
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _sanitize_excel_val(val: Any) -> Any:
    """Mencegah Excel formula injection jika string diawali =, +, -, @."""
    if isinstance(val, str):
        if val.startswith(("=", "+", "-", "@")):
            return "'" + val
    return val


def _style_header_row(ws, row_idx: int, col_count: int):
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="medium", color="0F172A"),
    )
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border


def _auto_fit_columns(ws, max_cols: int = 25):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            v_str = str(cell.value or "")
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws.column_dimensions[col_letter].width = max(11, min(max_len + 3, 40))


def build_xlsx_report(report_data: Dict[str, Any]) -> io.BytesIO:
    """Membangun workbook Excel XLSX multi-sheet komprehensif menggunakan openpyxl."""
    wb = openpyxl.Workbook()
    # Hapus sheet default
    wb.remove(wb.active)

    meta = report_data.get("metadata", {})
    ds = report_data.get("dataset", {})
    sel_sections = set(report_data.get("selected_sections", []))
    analysis_map = report_data.get("analysis_map", {})
    histories = report_data.get("histories", [])

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # 1. SHEET: RINGKASAN LAPORAN
    ws_summary = wb.create_sheet(title="Ringkasan")
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary.cell(row=2, column=2, value="LAPORAN HASIL ANALISIS STATISTIK").font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    ws_summary.cell(row=3, column=2, value="Statistika Pro — Applied Statistical Analysis System").font = Font(name="Calibri", size=11, color="64748B")

    summary_rows = [
        ("Judul Penelitian", meta.get("title") or "-"),
        ("Peneliti", meta.get("researcher") or "-"),
        ("Institusi / Jurusan", meta.get("institution") or "-"),
        ("Periode Penelitian", meta.get("period") or "-"),
        ("Nama Dataset", ds.get("name") or "-"),
        ("Ukuran Sampel (N)", len(ds.get("rows", []))),
        ("Jumlah Variabel", len(ds.get("variables", []))),
        ("Waktu Cetak", datetime.now().strftime("%d/%m/%Y %H:%M WIB")),
    ]
    if meta.get("notes"):
        summary_rows.append(("Catatan", meta.get("notes")))

    start_r = 5
    for idx, (lbl, val) in enumerate(summary_rows):
        r = start_r + idx
        c_lbl = ws_summary.cell(row=r, column=2, value=lbl)
        c_lbl.font = Font(name="Calibri", size=11, bold=True, color="334155")
        c_lbl.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        c_lbl.border = thin_border

        c_val = ws_summary.cell(row=r, column=3, value=_sanitize_excel_val(val))
        c_val.font = Font(name="Calibri", size=11)
        c_val.border = thin_border

    ws_summary.column_dimensions["B"].width = 24
    ws_summary.column_dimensions["C"].width = 45

    # 2. SHEET: KAMUS VARIABEL
    if "kamus_variabel" in sel_sections and ds.get("variables"):
        ws_vars = wb.create_sheet(title="Kamus Variabel")
        ws_vars.views.sheetView[0].showGridLines = True

        v_headers = ["No", "Nama Variabel", "Label Variabel", "Tipe", "Skala Ukur", "Missing Values"]
        for c_idx, h in enumerate(v_headers, start=1):
            ws_vars.cell(row=1, column=c_idx, value=h)
        _style_header_row(ws_vars, 1, len(v_headers))

        for r_idx, v in enumerate(ds.get("variables", []), start=2):
            missing_str = ", ".join(str(m) for m in v.get("missing", [])) or "None"
            row_vals = [
                r_idx - 1,
                v.get("name", ""),
                v.get("label", "") or "-",
                v.get("type", "numeric"),
                v.get("measure", "scale"),
                missing_str,
            ]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_vars.cell(row=r_idx, column=c_idx, value=_sanitize_excel_val(val))
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)

        _auto_fit_columns(ws_vars)

    # 3. SHEET: DATA AKTIF
    if ds.get("rows") and ds.get("variables"):
        ws_data = wb.create_sheet(title="Data")
        ws_data.views.sheetView[0].showGridLines = True

        var_names = [v.get("name") for v in ds.get("variables", [])]
        cols_to_write = ["id"] + [vn for vn in var_names if vn != "id"]

        for c_idx, col_name in enumerate(cols_to_write, start=1):
            ws_data.cell(row=1, column=c_idx, value=col_name)
        _style_header_row(ws_data, 1, len(cols_to_write))

        for r_idx, row_dict in enumerate(ds.get("rows", []), start=2):
            for c_idx, col_name in enumerate(cols_to_write, start=1):
                raw_v = row_dict.get(col_name)
                cell = ws_data.cell(row=r_idx, column=c_idx, value=_sanitize_excel_val(raw_v))
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)

        _auto_fit_columns(ws_data)

    # 4. SHEET RIWAYAT TRANSFORMASI
    if "riwayat_transformasi" in sel_sections and histories:
        ws_hist = wb.create_sheet(title="Transformasi")
        ws_hist.views.sheetView[0].showGridLines = True

        h_cols = ["Waktu (UTC)", "Operasi", "Keterangan", "Baris Terdampak", "Variabel Sumber", "Variabel Target"]
        for c_idx, h in enumerate(h_cols, start=1):
            ws_hist.cell(row=1, column=c_idx, value=h)
        _style_header_row(ws_hist, 1, len(h_cols))

        for r_idx, h_item in enumerate(histories, start=2):
            ts = h_item.get("timestamp", "").replace("T", " ")[:19]
            vals = [
                ts,
                h_item.get("operation_type", "-"),
                h_item.get("description", "-"),
                h_item.get("affected_rows", 0),
                ", ".join(h_item.get("source_vars", [])) or "-",
                ", ".join(h_item.get("target_vars", [])) or "-",
            ]
            for c_idx, val in enumerate(vals, start=1):
                cell = ws_hist.cell(row=r_idx, column=c_idx, value=_sanitize_excel_val(val))
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)

        _auto_fit_columns(ws_hist)

    # 5. SHEETS UNTUK ANALISIS STATISTIK
    section_mapping = [
        ("deskriptif", "descriptive", "Deskriptif"),
        ("reliabilitas", "reliability", "Reliabilitas"),
        ("validitas", "validity", "Validitas"),
        ("korelasi", "correlation", "Korelasi"),
        ("regresi", "regression", "Regresi"),
        ("diagnostik", "diagnostics", "Diagnostik"),
    ]

    for sec_key, anal_type, sheet_name in section_mapping:
        if sec_key not in sel_sections:
            continue
        outputs_list = analysis_map.get(anal_type, [])
        if not outputs_list:
            continue

        ws_stat = wb.create_sheet(title=sheet_name)
        ws_stat.views.sheetView[0].showGridLines = True

        curr_row = 1
        for out in outputs_list:
            # Header Judul Analisis
            ws_stat.cell(row=curr_row, column=1, value=out.title).font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
            curr_row += 1

            res = out.results
            if isinstance(res, dict):
                tables_to_draw = []
                if "tables" in res and isinstance(res["tables"], list):
                    tables_to_draw = res["tables"]
                elif "statistics" in res and isinstance(res["statistics"], list):
                    tables_to_draw = [{"title": "Ringkasan Statistik", "rows": res["statistics"]}]

                if tables_to_draw:
                    for t in tables_to_draw:
                        t_rows_raw = t.get("rows", [])
                        if not t_rows_raw:
                            continue
                        cols = list(t_rows_raw[0].keys())
                        for c_i, c_name in enumerate(cols, start=1):
                            ws_stat.cell(row=curr_row, column=c_i, value=str(c_name))
                        _style_header_row(ws_stat, curr_row, len(cols))
                        curr_row += 1

                        for r_item in t_rows_raw:
                            for c_i, c_name in enumerate(cols, start=1):
                                val = r_item.get(c_name)
                                cell = ws_stat.cell(row=curr_row, column=c_i, value=_sanitize_excel_val(val))
                                cell.border = thin_border
                                cell.font = Font(name="Calibri", size=10)
                            curr_row += 1
                        curr_row += 1
                else:
                    # Key value
                    for k, v in list(res.items())[:15]:
                        if isinstance(v, (int, float, str, bool)):
                            c0 = ws_stat.cell(row=curr_row, column=1, value=str(k))
                            c0.font = Font(name="Calibri", size=10, bold=True)
                            c1 = ws_stat.cell(row=curr_row, column=2, value=_sanitize_excel_val(v))
                            c1.font = Font(name="Calibri", size=10)
                            curr_row += 1

            if out.interpretation:
                c_lbl = ws_stat.cell(row=curr_row, column=1, value="Interpretasi:")
                c_lbl.font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
                c_txt = ws_stat.cell(row=curr_row, column=2, value=out.interpretation)
                c_txt.font = Font(name="Calibri", size=10, italic=True)
                curr_row += 2

            curr_row += 1

        _auto_fit_columns(ws_stat)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
